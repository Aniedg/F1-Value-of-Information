"""Editable body and tables; display equations are high-resolution images.

Run using the bundled document runtime after build_pdf_v2 and docx_payload.
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
payload=json.loads((ROOT/'tmp/docx-payload.json').read_text());doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',20),('Subtitle',12),('Heading 1',13),('Heading 2',11.5)]:
 style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
 style.paragraph_format.space_before=Pt(10);style.paragraph_format.space_after=Pt(6)
 if name in ['Heading 1','Heading 2']:style.font.bold=True
for style in doc.styles:
 for elem in list(style.element.iter(qn('w:pBdr'))):elem.getparent().remove(elem)
 for elem in style.element.iter(qn('w:rFonts')):
  for key in list(elem.attrib):
   if 'theme' in key.lower():del elem.attrib[key]
subtitle=False;table_count=0
captions=['Parameter grid for the protocol-locked synthetic study.','Equal-state mean costs and regret on the revised 2,187-state grid.','Robust forecasting on the three-event test group.','Adaptive and recalibrated static intervals on the final three-event group.','Reserved-event forecasts and conditional late-race decisions.']
for node in payload:
 t=node['type']
 if t=='h1':p=doc.add_paragraph(node['text'],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif t=='h2':
  if not subtitle:p=doc.add_paragraph(node['text'],'Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;subtitle=True
  else:
   if node['text']=='References':doc.add_page_break()
   doc.add_paragraph(node['text'],'Heading 1')
 elif t=='h3':doc.add_paragraph(node['text'],'Heading 2')
 elif t=='p':
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  for part in node['runs']:
   run=p.add_run(part['text']);run.bold=part['bold'];run.italic=part['italic'];run.font.subscript=part.get('subscript',False)
 elif t=='image':
  with Image.open(node['path']) as im:w,h=im.size
  width=min(6.5,w/240) if node['equation'] else 6.5
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.add_run().add_picture(node['path'],width=Inches(width))
  if not node['equation']:
   p.paragraph_format.keep_with_next=True
   cap=doc.add_paragraph(node['caption']);cap.paragraph_format.space_after=Pt(9)
   for r in cap.runs:r.font.size=Pt(9)
 elif t=='table':
  cap=doc.add_paragraph(f'Table {table_count+1}. '+captions[table_count]);cap.paragraph_format.keep_with_next=True
  for r in cap.runs:r.font.size=Pt(9)
  table_count+=1;table=doc.add_table(rows=0,cols=len(node['rows'][0]));table.style='Table Grid'
  for i,row in enumerate(node['rows']):
   cells=table.add_row().cells
   for cell,value in zip(cells,row):
    cell.text=value
    for p in cell.paragraphs:
     p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
     for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0);r.bold=i==0
   trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   if i==0:trpr.append(OxmlElement('w:tblHeader'))
  doc.add_paragraph().paragraph_format.space_after=Pt(0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
doc.core_properties.title='The Value of Information in Formula One Pit Strategy';doc.core_properties.author='';doc.core_properties.subject='Bayesian stopping and conditional historical replay';doc.core_properties.comments=''
out=ROOT/'output/docx';out.mkdir(exist_ok=True)
doc.save(out/'value-of-information-f1-paper-v2.docx')
print(out/'value-of-information-f1-paper-v2.docx')
